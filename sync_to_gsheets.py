"""
Синхронизация Excel → Google Sheets.
Дополнительно переносит примечания к столбцу V (частичные оплаты)
в колонку 'payment_parts'.
Обрезает служебные блоки внизу Excel.
Добавлена колонка 'Код хоз-ва'.
"""
import pandas as pd
import gspread
from google.oauth2.service_account import Credentials
from datetime import datetime
import sys
import openpyxl

# === НАСТРОЙКИ ===
EXCEL_PATH = r"C:\ЦЕНТР ИЗЖ\Х\!-Х12.1.xlsm"
SHEET_NAME = "ИСХОДНЫЕ ДАННЫЕ"
CREDENTIALS_FILE = r"C:\X12\service-account.json"
SPREADSHEET_ID = "1AWSwJECekzgfvbYlBsBk-Ws78hpNp5TC7VSPdvv0Nso"
WORKSHEET_NAME = "Data"

COL_V = 22  # столбец V (ПП, дата) — примечания про частичные оплаты

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive"
]

# Колонки, которые нужны data_loader.py
KEEP_COLUMNS = [
    'Код хоз-ва',           # ← НОВОЕ
    'Наименование хозяйства',
    'Область',
    'Район',
    'Менеджер',
    'Вид Исследования',
    'Счет, номер',
    'Счет, дата',
    'Счет, сумма',
    'ПП, сумма',
    'ПП, дата',
    'ПП, номер',
    'Срок оплаты, дни',
    'цифра в зав-ти от цвета в  столбце R',
    'Кол-во проб',
    'Сумма, BYN',
    'Номер Протокола',
    'Предприятие',
]

ALL_MANAGERS = [
    'Бабура С.', 'Крупенькина Е.', 'Чубарь Д.', 'Буян В.',
    'Пудакевич И.', 'Федук П.', 'Вашкевич А.', 'Шиянов В.',
    'Пломодьялов Д.', 'Сагайдак В.', 'Зварич Д.', 'Нестер А.',
    'Ребковец Е.', 'Бабайцев Е.',
]

DATE_COLUMNS = ['Счет, дата', 'ПП, дата']
EXTRA_COLUMN = 'payment_parts'


def load_comment_texts():
    print("Чтение примечаний из Excel...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False, read_only=False)
    ws = wb[SHEET_NAME]

    comments = {}
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=COL_V)
        if cell.comment and cell.comment.text:
            text = cell.comment.text.strip().replace('\n', ' ')
            if text:
                comments[row] = text

    wb.close()
    print(f"  Найдено {len(comments)} примечаний")
    return comments


def main():
    print(f"[{datetime.now().strftime('%H:%M:%S')}] Старт синхронизации...")

    print(f"Чтение Excel: {EXCEL_PATH}")
    try:
        df = pd.read_excel(EXCEL_PATH, sheet_name=SHEET_NAME, header=0)
    except Exception as e:
        print(f"ОШИБКА чтения Excel: {e}")
        sys.exit(1)

    print(f"  Загружено {len(df)} строк, {len(df.columns)} колонок")

    comments = load_comment_texts()

    available = [c for c in KEEP_COLUMNS if c in df.columns]
    missing = [c for c in KEEP_COLUMNS if c not in df.columns]
    if missing:
        print(f"ВНИМАНИЕ: не найдены колонки: {missing}")
    df = df[available].copy()
    print(f"  Оставлено {len(df.columns)} колонок")

    if 'Менеджер' in df.columns:
        mask = df['Менеджер'].isin(ALL_MANAGERS)
        if mask.any():
            last_valid = mask[mask].index[-1]
            before = len(df)
            df = df.loc[:last_valid].copy()
            after = len(df)
            if before != after:
                print(f"  Обрезано служебных строк: {before - after}")

    df[EXTRA_COLUMN] = [
        comments.get(int(i) + 2, '') if str(i).isdigit() else ''
        for i in df.index
    ]
    filled = (df[EXTRA_COLUMN] != '').sum()
    print(f"  Заполнено {EXTRA_COLUMN}: {filled} строк")

    for col in DATE_COLUMNS:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors='coerce')
            df[col] = df[col].dt.strftime('%Y-%m-%d').fillna('')

    df = df.fillna('')

    print("Авторизация в Google...")
    try:
        creds = Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=SCOPES)
        gc = gspread.authorize(creds)
    except Exception as e:
        print(f"ОШИБКА авторизации: {e}")
        sys.exit(1)

    try:
        sh = gc.open_by_key(SPREADSHEET_ID)
        worksheet = sh.worksheet(WORKSHEET_NAME)
    except Exception as e:
        print(f"ОШИБКА открытия таблицы: {e}")
        sys.exit(1)

    print(f"Отправка в Google Sheets (лист '{WORKSHEET_NAME}')...")
    try:
        worksheet.clear()
        data_to_write = [df.columns.tolist()] + df.astype(str).values.tolist()
        worksheet.update(data_to_write, value_input_option='RAW')
    except Exception as e:
        print(f"ОШИБКА записи: {e}")
        sys.exit(1)

    print(f"ГОТОВО! Записано {len(df)} строк, {len(df.columns)} колонок.")


if __name__ == "__main__":
    main()