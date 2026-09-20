"""
Парсер примечаний к ячейкам V (ПП, дата).
Формат примечания: СУММА_2 ДАТА_2 НОМЕР_2; СУММА_3 ДАТА_3 НОМЕР_3; ...
"""
import re
import openpyxl
from datetime import datetime

EXCEL_PATH = r"C:\ЦЕНТР ИЗЖ\Х\!-Х12.1.xlsm"
SHEET_NAME = "ИСХОДНЫЕ ДАННЫЕ"
COL_V = 22  # столбец V (ПП, дата)


def parse_comment(text):
    """
    Разбирает текст примечания.
    Возвращает список словарей:
      [{'amount': float, 'date': datetime, 'num': str}, ...]
    """
    if not text:
        return []

    result = []
    text = text.replace('\n', ' ').strip()

    # Разделяем по ';'
    parts = text.split(';')

    for part in parts:
        part = part.strip()
        if not part:
            continue

        # Ищем сумму (число), дату (ДД.ММ.ГГГГ) и опциональный номер
        # Формат: СУММА ДАТА [НОМЕР]
        tokens = part.split()
        if len(tokens) < 2:
            continue

        amount = None
        date = None
        num = None

        for i, tok in enumerate(tokens):
            # Сумма
            if amount is None:
                try:
                    amount = float(tok.replace(',', '.'))
                    continue
                except ValueError:
                    pass

            # Дата
            if date is None:
                try:
                    date = datetime.strptime(tok, '%d.%m.%Y')
                    continue
                except ValueError:
                    pass

            # Номер ПП (последний токен)
            if date is not None and num is None:
                num = tok

        if amount is not None and date is not None:
            result.append({
                'amount': amount,
                'date': date,
                'num': num
            })

    return result


def load_comments():
    """
    Возвращает словарь: {номер_строки_Excel: [части оплат из примечания]}
    Только для строк с примечанием в столбце V.
    """
    print("Чтение примечаний из Excel...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False, read_only=False)
    ws = wb[SHEET_NAME]

    comments = {}

    max_row = ws.max_row
    for row in range(2, max_row + 1):
        cell = ws.cell(row=row, column=COL_V)
        if cell.comment and cell.comment.text:
            parts = parse_comment(cell.comment.text)
            if parts:
                comments[row] = parts

    wb.close()
    print(f"✅ Найдено {len(comments)} примечаний к столбцу V")

    return comments


if __name__ == "__main__":
    comments = load_comments()
    print("\n--- Примеры ---")
    for i, (row, parts) in enumerate(comments.items()):
        print(f"\nСтрока {row}:")
        for p in parts:
            print(f"  {p['amount']} BYN, {p['date'].strftime('%d.%m.%Y')}, ПП {p['num']}")
        if i >= 5:
            break